import os

import openpyxl
import pytest
from openpyxl.styles import Border, Side

from etl.common import norm_label, normalize_code, parse_code, parse_value, strip_period
from etl.harmonize import describe_column, split_unit
from etl.headers import build_header_paths

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---- common ---------------------------------------------------------------------

@pytest.mark.parametrize('raw, expected', [
    (12.5, (12.5, None, None)),
    ('1,234', (1234.0, None, None)),
    ('-', (None, None, 'dash')),
    ('..', (None, None, 'not_available')),
    ('( 2.3)', (2.3, None, 'low_reliability')),
    ('הצפון', (None, 'הצפון', 'text')),
    (None, (None, None, 'blank')),
])
def test_parse_value(raw, expected):
    assert parse_value(raw) == expected


def test_codes():
    assert parse_code('0026') == 26
    assert normalize_code(26, 'מועצה מקומית') == 26          # ראש פינה
    assert normalize_code(26, 'מועצה אזורית') == 5526        # מטה יהודה (pre-2024 code)
    assert normalize_code(5501, 'מועצה אזורית') == 5501      # already 2024 scheme


def test_norm_label():
    assert norm_label('סך הכל  אוכלוסייה') == 'סה"כ אוכלוסייה'
    assert norm_label('כלל ארצי1') == 'כלל ארצי'
    assert norm_label('מעמד מוניציפאלי') == 'מעמד מוניציפלי'
    assert norm_label('יישובים 4,999-2,000 תושבים') == 'יישובים 2,000-4,999 תושבים'
    assert norm_label('בני 4-0') == 'בני 0-4'
    assert norm_label('תחבורה ימית 2013 -אחוז') == 'תחבורה ימית 2013 - אחוז'


def test_strip_period():
    assert strip_period('סה"כ אוכלוסייה בסוף 2014 (אלפים)') == 'סה"כ אוכלוסייה בסוף השנה (אלפים)'
    assert strip_period('בתי ספר - סה"כ תשע"ה 2014/15') == 'בתי ספר - סה"כ'


# ---- harmonize ------------------------------------------------------------------

def test_split_unit():
    assert split_unit('צריכת מים (אלפי מ"ק)') == ('צריכת מים', ['אלפי מ"ק'])
    assert split_unit('חיובי ארנונה לפי סוג נכס - באלפי ש"ח') == ('חיובי ארנונה לפי סוג נכס', ['אלפי ש"ח'])


def test_describe_column_old_and_new_align():
    old = describe_column('אחוז בני 4-0 באוכלוסייה בסוף 2014', '', '', 2014)
    new = describe_column('דמוגרפיה › אחוז באוכלוסייה בסוף השנה › בני 4-0', '', '', 2024)
    assert old['data_period'] == '2014' and old['period_source'] == 'label'
    assert new['data_period'] == '2024' and new['period_source'] == 'file_default'
    assert sorted(old['flat_no_top'].split()) == sorted(new['flat_no_top'].split())


def test_describe_column_flat_label_keeps_first_level():
    d = describe_column('גמר בנייה: שטח (אלפי מ"ר) - סה"כ 2015', '', '', 2015)
    assert d['flat_no_top'] == d['flat']          # no category to drop in flat labels


def test_population_unit_variant():
    a = describe_column('דמוגרפיה › סה"כ אוכלוסייה בסוף השנה (אלפים)', '', '', 2017)
    b = describe_column('דמוגרפיה › סה"כ אוכלוסייה בסוף השנה', '', '', 2018)
    assert a['key'] != b['key'] and a['base_key'] == b['base_key']
    assert a['unit_hint'] == 'אלפים' and b['unit_hint'] == ''


# ---- headers --------------------------------------------------------------------

def _ws(rows, borders=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    thin = Side(style='thin')
    for (r, c, side) in borders:                     # 0-based row/col
        cell = ws.cell(row=r + 1, column=c + 1)
        cell.border = Border(left=thin if side == 'L' else None, right=thin if side == 'R' else None)
    return ws


def test_border_block_spreads_centred_label():
    # group label written in the middle (col 4) of a block bordered at cols 2..5
    rows = [['שם', 'סמל', None, None, 'קבוצה', None],
            ['שם', 'סמל', 'א', 'ב', 'ג', 'ד']]
    ws = _ws(rows, borders=[(0, 1, 'R'), (0, 5, 'R')])
    paths = build_header_paths(ws, [0, 1], None, None, 6)
    assert [p['levels'] for p in paths[2:]] == [['קבוצה', x] for x in 'אבגד']


def test_leaf_row_does_not_leak():
    rows = [['קטגוריה', None, None, None],
            [None, None, None, None],
            ['שם', 'סמל', 'מדד א', 'מדד ב'],
            [None, None, 'גברים', None]]
    ws = _ws(rows)
    paths = build_header_paths(ws, [0, 1, 2, 3], 1, 2, 4, leaf_rows=[3])
    assert paths[2]['levels'] == ['קטגוריה', 'מדד א', 'גברים']
    assert paths[3]['levels'] == ['קטגוריה', 'מדד ב']


# ---- real files (skipped when the CBS files are absent) -------------------------

needs_files = pytest.mark.skipif(not os.path.exists(os.path.join(ROOT, 'p_libud_24.xlsx')),
                                 reason='CBS source files not present')


@needs_files
def test_mislabeled_2019_file_detected():
    from etl.validate import detect_title_year, numeric_fingerprint
    assert detect_title_year(os.path.join(ROOT, 'p_libud_2019.xlsx')) == 2022
    assert numeric_fingerprint(os.path.join(ROOT, 'p_libud_2019.xlsx')) == \
        numeric_fingerprint(os.path.join(ROOT, 'p_libud_22.xlsx'))


@needs_files
def test_2017_rc_codes_normalized_and_collision_resolved():
    from etl.extract import extract_year, load_config
    cfg = load_config(os.path.join(ROOT, 'etl', 'config', 'sources.yaml'))
    r = extract_year(2017, cfg['years'][2017], ROOT, cfg.get('name_aliases'))
    reg = r['registry'].set_index('muni_code')
    assert reg.loc[5501, 'name'] == 'הגליל העליון'
    assert reg.loc[26, 'name'] == 'ראש פינה'
    assert reg.loc[5526, 'name'] == 'מטה יהודה'
    assert reg.index.is_unique
    # the budget sheet has no status column: codes resolved through the registry
    budget = r['cells'][r['cells']['sheet'] == 'נתוני תקציב']
    assert {5501, 26, 5526} <= set(budget['muni_code'].dropna().astype(int))
